import streamlit as st
import openpyxl
import io
import requests
import plotly.graph_objects as go
from PIL import Image, ImageDraw, ImageFont

st.set_page_config(
    page_title="GESTOR PAGOS VALLE HERMOSO",
    page_icon="🏡",
    layout="centered",
    initial_sidebar_state="collapsed"
)

# Estilos CSS compactos para móvil
st.markdown("""
<style>
    .block-container {
        padding-top: 1rem !important;
        padding-bottom: 4rem !important;
        padding-left: 0.6rem !important;
        padding-right: 0.6rem !important;
        max-width: 520px !important;
    }
    .app-title {
        font-size: 1.25rem;
        font-weight: 800;
        color: #1A365D;
        text-align: center;
        margin-bottom: 0.8rem;
    }
    /* Reducción de métricas nativas */
    [data-testid="stMetric"] {
        background-color: #F8FAFC;
        border: 1px solid #E2E8F0;
        border-radius: 8px;
        padding: 6px 10px;
    }
    [data-testid="stMetricLabel"] {
        font-size: 0.72rem !important;
        color: #64748B !important;
        font-weight: 600;
    }
    [data-testid="stMetricValue"] {
        font-size: 1.05rem !important;
        font-weight: 700 !important;
        color: #1E293B !important;
    }
    /* Contenedor compacto de las 36 cuotas */
    .cuotas-grid {
        display: grid;
        grid-template-columns: repeat(6, 1fr);
        gap: 4px;
        margin-top: 6px;
        margin-bottom: 12px;
    }
    .badge-c {
        font-size: 0.72rem;
        font-weight: 700;
        text-align: center;
        padding: 5px 0;
        border-radius: 5px;
        color: #ffffff;
    }
    .c-verde { background-color: #00C853; }
    .c-naranja { background-color: #FF9100; }
    .c-rojo { background-color: #E53935; }
    .c-blanco { background-color: #E2E8F0; color: #475569 !important; }
</style>
""", unsafe_allow_html=True)

SHEET_ID = "1CdRBWSW9QDh63s7-8lnZvGu8rwz65z2duPnk2Nqfzqc"

@st.cache_data(ttl=15)
def descargar_workbook():
    url = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"
    res = requests.get(url)
    if res.status_code != 200:
        return None
    wb = openpyxl.load_workbook(io.BytesIO(res.content), data_only=True)
    return wb

st.markdown('<div class="app-title">🏡 GESTOR PAGOS VALLE HERMOSO</div>', unsafe_allow_html=True)

if st.button("🔄 Refrescar Datos", use_container_width=True):
    st.cache_data.clear()
    st.rerun()

wb = descargar_workbook()
if not wb or "CT" not in wb.sheetnames:
    st.error("No se pudo conectar con la hoja de cálculo o no existe la pestaña 'CT'.")
    st.stop()

ws_ct = wb["CT"]
ws_ventas = wb["VENTAS"] if "VENTAS" in wb.sheetnames else None

def get_color_estado(cell):
    if not cell:
        return "BLANCO"

    val_str = str(cell.value or "").strip().upper()
    tiene_contenido = bool(val_str and val_str not in ["NONE", "NAN", ""])

    rgb_str = ""
    indexed_val = None

    if cell.fill and (cell.fill.fill_type or cell.fill.start_color or cell.fill.fgColor):
        color_obj = getattr(cell.fill, 'fgColor', None) or getattr(cell.fill, 'start_color', None)
        if color_obj:
            v_rgb = getattr(color_obj, 'rgb', None) or getattr(color_obj, 'value', None)
            if v_rgb and isinstance(v_rgb, str):
                rgb_str = v_rgb.upper().strip()
                if len(rgb_str) == 8:
                    rgb_str = rgb_str[2:]
            indexed_val = getattr(color_obj, 'indexed', None) or getattr(color_obj, 'index', None)

    verdes_rgb = [
        "00FF00", "57BB8A", "6AA84F", "00E676", "38761D", "85E89D", 
        "B7E1CD", "274E13", "81C784", "4CAF50", "2E7D32", "00C853", 
        "34A853", "137333", "0F9D58", "A8DAB5", "D9EAD3"
    ]
    if any(v in rgb_str for v in verdes_rgb) or indexed_val in [11, 17, 42, 43]:
        return "VERDE"

    naranjas_rgb = [
        "FF9900", "FF6D01", "FFA500", "F6B26B", "F9AB00", "E69138", 
        "FB8C00", "FF5722", "FF7043", "FFB74D", "FF9800", "E55100", 
        "B45F06", "783F04", "F9CB9C", "FCE5CD", "FF8A65", "FFAB40"
    ]
    if any(n in rgb_str for n in naranjas_rgb) or indexed_val in [51, 52, 53, 54]:
        return "NARANJA"

    rojos_rgb = ["FF0000", "CC0000", "E06666", "EA4335", "990000", "E57373", "F44336", "C62828"]
    if any(r in rgb_str for r in rojos_rgb) or indexed_val in [10, 16]:
        return "ROJO"

    if tiene_contenido:
        if any(w in val_str for w in ["BAJA", "PERDIDO", "CANCELADO"]):
            return "ROJO"
        return "NARANJA"

    return "BLANCO"

def clean_number(val):
    if val is None:
        return 0.0
    s = str(val).replace("S/.", "").replace("S/", "").replace(",", "").replace(" ", "").strip()
    try:
        return float(s)
    except:
        return 0.0

def normalizar(txt):
    if not txt:
        return ""
    return str(txt).replace(" ", "").replace("-", "").replace("_", "").upper().strip()

PALABRAS_IGNORAR = {
    "ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO",
    "JULIO", "AGOSTO", "SETIEMBRE", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE",
    "NOMBRE", "NOMBRES", "NOMBRRE", "CLIENTE", "ID", "TOTAL", "SUBTOTAL"
}

fila_encabezado_idx = 1
col_cuota_1 = 13

for r in range(1, 10):
    for c in range(8, 20):
        val = str(ws_ct.cell(row=r, column=c).value).strip()
        if val == "1":
            fila_encabezado_idx = r
            col_cuota_1 = c
            break
    if col_cuota_1 != 13:
        break

registros_lotes = []
for r in range(fila_encabezado_idx + 1, ws_ct.max_row + 1):
    val_nom = ws_ct.cell(row=r, column=1).value
    if val_nom is None:
        continue
    nom_str = str(val_nom).strip()
    if not nom_str or nom_str.upper() in PALABRAS_IGNORAR:
        continue
        
    etapa = str(ws_ct.cell(row=r, column=2).value or "").strip()
    lote = str(ws_ct.cell(row=r, column=3).value or "").strip()
    if not etapa and not lote:
        continue

    label = f"{nom_str} — {etapa} Lote {lote}"
    registros_lotes.append({
        "id_compuesto": label,
        "nombre": nom_str,
        "etapa": etapa,
        "lote": lote,
        "row_idx": r
    })

if not registros_lotes:
    st.warning("No se encontraron clientes.")
    st.stop()

opciones_busqueda = [item["id_compuesto"] for item in registros_lotes]

cliente_sel_label = st.selectbox(
    "🔍 Buscar por Cliente o Lote:",
    options=opciones_busqueda,
    index=None,
    placeholder="Escribe el nombre o lote..."
)

def generar_imagen_reporte(nombre, etapa, lote, m2, modalidad, valor_total, inicial, total_pagado, saldo, valor_cuota, pagadas, vencidas, cuotas_estados, es_contado):
    """Crea una tarjeta de imagen PNG en alta definición lista para compartir por WhatsApp."""
    w, h = 600, 780
    img = Image.new("RGB", (w, h), "#FFFFFF")
    draw = ImageDraw.Draw(img)

    # Encabezado azul oscuro
    draw.rectangle([(0, 0), (w, 95)], fill="#1A365D")
    draw.text((25, 20), "VALLE HERMOSO INMOBILIARIA", fill="#FFFFFF")
    draw.text((25, 48), "ESTADO DE CUENTA DE LOTE", fill="#90CDF4")

    # Datos del cliente
    draw.text((25, 110), f"CLIENTE: {nombre.upper()}", fill="#1E293B")
    draw.text((25, 135), f"UBICACIÓN: {etapa} — Lote {lote} ({m2} m²)", fill="#64748B")
    draw.text((420, 110), f"MODALIDAD: {modalidad}", fill="#2B6CB0")
    draw.line([(25, 165), (w - 25, 165)], fill="#E2E8F0", width=2)

    # Cuadrícula de valores
    datos_box = [
        ("VALOR TOTAL", f"S/. {valor_total:,.2f}"),
        ("INICIAL", f"S/. {inicial:,.2f}"),
        ("TOTAL PAGADO", f"S/. {total_pagado:,.2f}"),
        ("SALDO PENDIENTE", f"S/. {saldo:,.2f}"),
        ("VALOR CUOTA", f"S/. {valor_cuota:,.2f}"),
        ("CUOTAS PAGADAS", f"{pagadas} de 36 ({vencidas} venc)" if vencidas > 0 else f"{pagadas} de 36")
    ]

    for idx, (label, valor) in enumerate(datos_box):
        x = 25 if idx % 2 == 0 else 310
        y = 185 + (idx // 2) * 65
        draw.rectangle([(x, y), (x + 265, y + 55)], fill="#F8FAFC", outline="#E2E8F0", width=1)
        draw.text((x + 12, y + 8), label, fill="#64748B")
        color_val = "#2E7D32" if "PAGADO" in label else ("#DD6B20" if "SALDO" in label else "#1E293B")
        draw.text((x + 12, y + 26), valor, fill=color_val)

    # Matriz visual de 36 cuotas
    draw.text((25, 395), "MATRIZ DE 36 CUOTAS (Verde: Pagada | Naranja: Vencida)", fill="#1E293B")
    draw.line([(25, 420), (w - 25, 420)], fill="#E2E8F0", width=1)

    c_box_w, c_box_h = 85, 32
    start_x, start_y = 25, 435

    for i in range(1, 37):
        col = (i - 1) % 6
        fil = (i - 1) // 6
        x = start_x + col * (c_box_w + 8)
        y = start_y + fil * (c_box_h + 8)

        est = cuotas_estados[i - 1]
        fill_color = "#00C853" if (es_contado or est == "VERDE") else ("#FF9100" if est == "NARANJA" else ("#E53935" if est == "ROJO" else "#E2E8F0"))
        draw.rectangle([(x, y), (x + c_box_w, y + c_box_h)], fill=fill_color)
        txt_color = "#FFFFFF" if fill_color != "#E2E8F0" else "#475569"
        draw.text((x + 30, y + 8), f"C{i}", fill=txt_color)

    # Pie de reporte
    draw.line([(25, 715), (w - 25, 715)], fill="#E2E8F0", width=1)
    draw.text((25, 730), "Reporte oficial generado desde Gestor Pagos Valle Hermoso", fill="#94A3B8")

    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()

if cliente_sel_label:
    data_sel = next(item for item in registros_lotes if item["id_compuesto"] == cliente_sel_label)
    row_idx = data_sel["row_idx"]
    nom_cliente = data_sel["nombre"]
    etapa = data_sel["etapa"]
    lote = data_sel["lote"]

    m2 = str(ws_ct.cell(row=row_idx, column=4).value or "-").strip()
    valor_cuota = clean_number(ws_ct.cell(row=row_idx, column=9).value)

    valor_total = clean_number(ws_ct.cell(row=row_idx, column=5).value)
    monto_inicial = clean_number(ws_ct.cell(row=row_idx, column=6).value)
    modalidad = "FINANCIADO"

    if ws_ventas:
        for r_v in range(2, ws_ventas.max_row + 1):
            nom_v = str(ws_ventas.cell(row=r_v, column=2).value or "").strip()
            lote_v = str(ws_ventas.cell(row=r_v, column=5).value or "").strip()
            
            coincide_nom = normalizar(nom_v) == normalizar(nom_cliente)
            coincide_lote = normalizar(lote_v) == normalizar(lote) if lote and lote_v else True
            
            if coincide_nom and coincide_lote:
                v_total = clean_number(ws_ventas.cell(row=r_v, column=7).value)
                v_inic = clean_number(ws_ventas.cell(row=r_v, column=12).value)
                mod_v = str(ws_ventas.cell(row=r_v, column=15).value or "").strip().upper()
                
                if v_total > 0: valor_total = v_total
                if v_inic > 0: monto_inicial = v_inic
                if mod_v: modalidad = mod_v
                break

    cuotas_estados = []
    total_verdes = 0
    total_naranjas = 0
    total_rojas = 0

    for i in range(36):
        c_idx = col_cuota_1 + i
        cell = ws_ct.cell(row=row_idx, column=c_idx)
        est = get_color_estado(cell)
        
        if est == "VERDE":
            total_verdes += 1
            cuotas_estados.append("VERDE")
        elif est == "NARANJA":
            total_naranjas += 1
            cuotas_estados.append("NARANJA")
        elif est == "ROJO":
            total_rojas += 1
            cuotas_estados.append("ROJO")
        else:
            cuotas_estados.append("BLANCO")

    es_contado = "CONTADO" in modalidad
    if es_contado:
        total_pagado = valor_total
        saldo_pendiente = 0.0
    else:
        total_pagado = monto_inicial + (total_verdes * valor_cuota)
        saldo_pendiente = max(0.0, valor_total - total_pagado)

    # Tarjeta de Datos Nativos (Sin HTML que falle)
    with st.container(border=True):
        col_t1, col_t2 = st.columns([3, 1])
        with col_t1:
            st.subheader(nom_cliente)
            st.caption(f"📍 {etapa} — Lote {lote} ({m2} m²)")
        with col_t2:
            st.info(modalidad)

        # Grilla de métricas en 2 columnas para celular
        m_c1, m_c2 = st.columns(2)
        m_c1.metric("Valor Total", f"S/. {valor_total:,.2f}")
        m_c2.metric("Monto Inicial", f"S/. {monto_inicial:,.2f}")

        m_c3, m_c4 = st.columns(2)
        m_c3.metric("Total Pagado", f"S/. {total_pagado:,.2f}")
        m_c4.metric("Saldo Pendiente", f"S/. {saldo_pendiente:,.2f}")

        m_c5, m_c6 = st.columns(2)
        m_c5.metric("Valor Cuota", f"S/. {valor_cuota:,.2f}")
        m_c6.metric("Cuotas Pagadas", f"{total_verdes} de 36", f"{total_naranjas} vencidas" if total_naranjas > 0 else None)

        # Cuadrícula compacta de 36 cuotas (HTML en una sola línea)
        st.markdown("**Matriz de 36 Cuotas**")
        st.caption("🟢 Verde: Pagada | 🟠 Naranja: Vencida | ⚪ Neutro: Pendiente")
        badges = "".join([f'<div class="badge-c {"c-verde" if (es_contado or cuotas_estados[i-1]=="VERDE") else ("c-naranja" if cuotas_estados[i-1]=="NARANJA" else ("c-rojo" if cuotas_estados[i-1]=="ROJO" else "c-blanco"))}">C{i}</div>' for i in range(1, 37)])
        st.markdown(f'<div class="cuotas-grid">{badges}</div>', unsafe_allow_html=True)

    # Gráfico de dona compacto
    fig = go.Figure(data=[go.Pie(
        labels=['Total Pagado', 'Saldo Pendiente'],
        values=[total_pagado, saldo_pendiente],
        hole=.55,
        marker_colors=['#00E676', '#FFA500'] if not es_contado else ['#00E676', '#CCCCCC'],
        textinfo='percent'
    )])
    fig.update_layout(
        height=200,
        margin=dict(t=5, b=5, l=5, r=5),
        showlegend=True,
        legend=dict(orientation="h", yanchor="bottom", y=-0.25, xanchor="center", x=0.5, font=dict(size=11))
    )
    st.plotly_chart(fig, use_container_width=True, config={'displayModeBar': False})

    # Botón para generar y descargar la imagen del reporte
    img_bytes = generar_imagen_reporte(
        nombre=nom_cliente,
        etapa=etapa,
        lote=lote,
        m2=m2,
        modalidad=modalidad,
        valor_total=valor_total,
        inicial=monto_inicial,
        total_pagado=total_pagado,
        saldo=saldo_pendiente,
        valor_cuota=valor_cuota,
        pagadas=total_verdes,
        vencidas=total_naranjas,
        cuotas_estados=cuotas_estados,
        es_contado=es_contado
    )

    st.download_button(
        label="📸 Descargar / Compartir Ficha (Imagen PNG)",
        data=img_bytes,
        file_name=f"Estado_Cuenta_{nom_cliente.replace(' ', '_')}.png",
        mime="image/png",
        use_container_width=True
    )
